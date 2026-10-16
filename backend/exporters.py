"""File exporters consume one saved, validated course snapshot."""

from copy import deepcopy
import csv
from datetime import date
from io import BytesIO, StringIO
import math
from pathlib import Path
import re
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt

from backend.models import ProjectRecord
from backend.validation import HOUR_FIELDS


TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"


def _paragraph(doc, text, heading=None):
    paragraph = doc.add_heading(text, level=heading) if heading else doc.add_paragraph(text)
    return paragraph


def _list(doc, values):
    for index, value in enumerate(values, 1):
        _paragraph(doc, f"{index}. {value}")


def _table(doc, headers, rows, fractions=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    fractions = fractions or [1 / len(headers)] * len(headers)
    for column, fraction in zip(table.columns, fractions):
        column.width = int(width * fraction)
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = int(width * fractions[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    edge.set(qn(f"w:{key}"), value)
                borders.append(edge)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:w"), "65")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            properties.append(margins)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7E6E6")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                if len(headers) > 5 and column_index not in {1, len(headers)-1}:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9 if len(headers) > 5 else 10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def _new_document(course):
    source = Document(TEMPLATE_DIR / f"{course.template_type}.docx")
    metadata = [deepcopy(p._p) for p in source.paragraphs[:6] if p.text.strip() and set(p.text.strip()) != {"_"}]
    section_patterns = [deepcopy(section._sectPr) for section in source.sections]
    body = source._element.body
    for child in list(body):
        body.remove(child)
    body.append(deepcopy(section_patterns[0]))
    for block in metadata:
        body.insert(len(body)-1, block)
    for role in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        source.styles[role].font.color.rgb = RGBColor(0, 0, 0)
    source.core_properties.title = course.title
    source.core_properties.author = "Академия Информационных Технологий"
    source.core_properties.last_modified_by = ""
    return source, section_patterns


def _section(doc, pattern):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    # Use the source's exact geometry while preserving generated section relationships.
    for tag in ("pgSz", "pgMar", "cols"):
        existing = section._sectPr.find(qn(f"w:{tag}"))
        if existing is not None:
            section._sectPr.remove(existing)
        value = pattern.find(qn(f"w:{tag}"))
        if value is not None:
            section._sectPr.append(deepcopy(value))


def _cover(doc, record):
    course = record.course
    paragraph = doc.add_paragraph("УТВЕРЖДЕНО")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph = doc.add_paragraph('Приказом № ______ от «___» __________ 20__ г.')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(36)
    title = "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА ПОВЫШЕНИЯ КВАЛИФИКАЦИИ" if course.template_type == "dpp_pk" else "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА"
    for value in (title, f"«{course.title}»"):
        paragraph = doc.add_paragraph(value, style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
    if record.execution_mode == "demo":
        paragraph = doc.add_paragraph("Демонстрационный черновик для проверки работы сервиса")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph(f"Магас {date.today().year}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def _competencies(doc, course):
    for comp in course.competencies:
        _paragraph(doc, f"{comp.id}. {comp.name}")
        if course.template_type == "dpp_pk":
            _paragraph(doc, f"Профстандарт: {comp.profstandard_code}. Трудовая функция: {comp.labor_function_code}. Требует проверки методистом.")
        for label, values in (("Знания", comp.knowledge), ("Умения", comp.skills), ("Практический опыт и навыки", comp.practical_experience)):
            _paragraph(doc, label)
            _list(doc, values)


def _hours(doc, course, detailed):
    headers = ["№", "Раздел или тема", "Всего", "ЛК", "ПР", "ЛБ", "СРС", "ИА", "Контроль"]
    rows = []
    for mod in course.modules:
        rows.append([mod.id, mod.title, mod.total_hours, *(getattr(mod, name) for name in HOUR_FIELDS), 0, mod.attestation_form])
        if detailed:
            for topic in mod.topics:
                rows.append([topic.id, topic.title, topic.total_hours, *(getattr(topic, name) for name in HOUR_FIELDS), 0, topic.control_form])
    rows.append(["", "Итоговая аттестация", course.ia_hours, 0, 0, 0, 0, course.ia_hours, course.ia_form])
    rows.append(["", "Итого", course.total_hours, *(sum(getattr(mod, name) for mod in course.modules) for name in HOUR_FIELDS), course.ia_hours, ""])
    _table(doc, headers, rows, [.04, .38, .07, .055, .055, .055, .06, .055, .23])
    _paragraph(doc, "ЛК — лекции, ПР — практические занятия, ЛБ — лабораторные работы, СРС — самостоятельная работа, ИА — итоговая аттестация. Общий объём включает ИА один раз. Строки модулей являются итогами входящих в них тем.")


def _calendar(doc, course):
    _paragraph(doc, f"Срок реализации: {course.duration_days} учебных дней. Даты начала обучения и расписание заполняются для учебной группы.")
    # One consistent derived schedule; it never changes the approved course hours.
    capacity = max(1, math.ceil(course.total_hours / max(1, course.duration_days)))
    activities = [(f"{topic.id}. {topic.title}", topic.total_hours) for mod in course.modules for topic in mod.topics]
    activities.append((f"Итоговая аттестация: {course.ia_form}", course.ia_hours))
    rows = []
    day, remaining = 1, capacity
    for title, hours in activities:
        while hours:
            amount = min(remaining, hours)
            rows.append([day, title, amount])
            hours -= amount
            remaining -= amount
            if remaining == 0:
                day += 1
                remaining = capacity
    _table(doc, ["Учебный день", "Содержание", "Часы"], rows, [.14, .74, .12])


def _working_program(doc, course):
    for mod in course.modules:
        _paragraph(doc, mod.title, 3)
        for topic in mod.topics:
            _paragraph(doc, f"{topic.id}. {topic.title}", 3)
            _paragraph(doc, f"Объём: {topic.total_hours} ч. Связанные результаты: {', '.join(topic.linked_competencies)}. Контроль: {topic.control_form}.")
            _list(doc, topic.content_bullets)
            _paragraph(doc, topic.practice_description or "Описание практической или самостоятельной работы требует заполнения.")


def _assessment(doc, course):
    if course.template_type == "dpp_pk":
        _paragraph(doc, "2.5.1. Формы текущей, промежуточной и итоговой аттестации", 3)
    _paragraph(doc, f"Текущий контроль проводится в формах, указанных в учебном плане. Итоговая аттестация: {course.ia_form}, {course.ia_hours} ч.")
    _paragraph(doc, "Оценочные материалы и критерии семинаров приведены в приложениях. Требования к итоговой работе и критерии её защиты заполняются методистом, если не представлены в материалах ML-навыка.")
    if course.template_type == "dpp_pk":
        _paragraph(doc, "2.5.2. Оценочные материалы", 3)
    rows = []
    for comp in course.competencies:
        topics = [t.id for mod in course.modules for t in mod.topics if comp.id in t.linked_competencies]
        tests = [q.id for q in course.tests if q.linked_competency_id == comp.id]
        seminars = [s.id for s in course.seminars if s.linked_topic_id in topics]
        rows.append([comp.id, ", ".join(topics), ", ".join(tests), ", ".join(seminars)])
    _table(doc, ["Компетенция", "Темы", "Вопросы", "Задания"], rows, [.13, .17, .35, .35])
    if course.template_type == "dpp_pk":
        _paragraph(doc, "2.5.3. Анализ удовлетворённости требований слушателей", 3)
        _paragraph(doc, "Порядок анкетирования слушателей до и после обучения и ответственный за анализ результатов: ____________________")


def _materials(doc, course):
    _paragraph(doc, "Комплект включает рабочие материалы тем, задания семинаров, тесты с ключами и презентации. Слайды экспортируются отдельным файлом PPTX.")
    for mod in course.modules:
        for topic in mod.topics:
            if topic.lecture_hours:
                _paragraph(doc, f"План лекции по теме {topic.id}. {topic.title}", 3)
                _list(doc, topic.content_bullets or ["План лекции требует заполнения."])


def _conditions(doc, course, general):
    _paragraph(doc, "3. Организационно-педагогические условия реализации программы", 1)
    _paragraph(doc, "3.1. Материально-технические условия реализации программы", 2)
    software = [lit.reference for lit in course.literature if lit.category == "software"]
    _list(doc, software or ["Оборудование, программное обеспечение и доступы: ____________________"])
    _paragraph(doc, "Конкретная площадка проведения и состав оборудования: ____________________")
    _paragraph(doc, "3.2. Учебно-методическое и информационное обеспечение программы", 2)
    _paragraph(doc, "Рабочие программы тем, тестовые вопросы, семинарские задания, планы лекций и презентации входят в комплект программы.")
    if general:
        _paragraph(doc, "3.3. Литература", 2)
    categories = {"normative": "Нормативные документы", "main": "Основная литература", "additional": "Дополнительная литература", "internet": "Интернет-ресурсы", "software": "Программное обеспечение"}
    for category_index, (category, title) in enumerate(categories.items(), 1):
        items = [item for item in course.literature if item.category == category]
        if items:
            _paragraph(doc, title if general else f"3.2.{category_index}. {title}", 3)
            for index, item in enumerate(items, 1):
                value = f"{index}. {item.reference}"
                if item.url:
                    value += f" {item.url}"
                if item.needs_verification and "Требует проверки" not in value:
                    value += " [Требует проверки методистом]"
                _paragraph(doc, value)
    if not general:
        _paragraph(doc, "3.3. Организация образовательного процесса", 2)
        _paragraph(doc, f"Форма обучения: {course.education_form}. Расписание и состав учебной группы утверждаются организацией отдельно.")
    _paragraph(doc, "3.4. Кадровое обеспечение программы", 2)
    _table(doc, ["Модуль", "ФИО преподавателя", "Учёная степень или звание", "Должность"], [[mod.title, "________________", "________________", "________________"] for mod in course.modules], [.4, .25, .18, .17])
    _paragraph(doc, "Руководитель: ____________________    Подпись: ____________________")


def _appendices(doc, course):
    _paragraph(doc, "Приложение 1 Тестовые вопросы и ключи", 1)
    for question in course.tests:
        _paragraph(doc, f"{question.id}. {question.question}")
        _paragraph(doc, f"Тема: {question.linked_topic_id}. Компетенция: {question.linked_competency_id}.")
        for index, option in enumerate(question.options, 1):
            _paragraph(doc, f"{index}) {option}")
        _paragraph(doc, f"Ключ: {question.correct_answer}")
    _paragraph(doc, "Приложение 2 Семинарские задания", 1)
    for task in course.seminars:
        _paragraph(doc, f"{task.id}. {task.title}", 2)
        _paragraph(doc, f"Тема: {task.linked_topic_id}. {task.scenario}")
        _paragraph(doc, "Результат работы")
        _list(doc, task.deliverables)
        _paragraph(doc, "Критерии оценки")
        _list(doc, task.evaluation_criteria)


def export_docx(record: ProjectRecord) -> bytes:
    course = record.course
    doc, patterns = _new_document(course)
    _cover(doc, record)
    if course.template_type == "dpp_pk":
        _paragraph(doc, "1. Общая характеристика дополнительной профессиональной программы", 1)
        for heading, value in [("1.1. Аннотация программы", course.annotation), ("1.2. Цель реализации программы", course.goal)]:
            _paragraph(doc, heading, 2)
            _paragraph(doc, value)
        _list(doc, course.tasks)
        _paragraph(doc, "1.3. Планируемые результаты обучения", 2)
        _competencies(doc, course)
        _paragraph(doc, "1.4. Основания разработки программы", 2)
        _list(doc, [item.reference for item in course.literature if item.category == "normative"] or ["Нормативные основания требуют заполнения и проверки."])
        _paragraph(doc, "1.5. Перечень совершенствуемых профессиональных компетенций", 2)
        _list(doc, [f"{comp.id}. {comp.name}" for comp in course.competencies])
        for heading, value in [("1.6. Требования к слушателям", course.target_audience), ("1.7. Форма обучения", course.education_form), ("1.8. Объём программы", f"{course.total_hours} академических часов, включая {course.ia_hours} ч. итоговой аттестации. Срок: {course.duration_days} учебных дней."), ("1.9. Выдаваемый документ", course.issued_document)]:
            _paragraph(doc, heading, 2)
            _paragraph(doc, value)
        _paragraph(doc, "2. Содержание программы", 1)
        _paragraph(doc, "2.1. Учебный план", 2)
        _hours(doc, course, False)
        _paragraph(doc, "2.2. Учебно-тематический план", 2)
        _hours(doc, course, True)
        _paragraph(doc, "2.3. Календарный учебный график", 2)
        _calendar(doc, course)
        _paragraph(doc, "2.4. Рабочие программы модулей", 2)
        _working_program(doc, course)
        _paragraph(doc, "2.5. Оценка качества освоения программы", 2)
        _assessment(doc, course)
        _materials(doc, course)
    else:
        _paragraph(doc, "Пояснительная записка", 1)
        _paragraph(doc, course.annotation)
        for heading, value in [("1.1. Актуальность программы", course.annotation), ("1.2. Категории обучающихся", course.target_audience), ("1.3. Цель и задачи программы", course.goal)]:
            _paragraph(doc, heading, 2)
            _paragraph(doc, value)
        _list(doc, course.tasks)
        _paragraph(doc, "1.4. Сроки реализации", 2)
        _paragraph(doc, f"Объём: {course.total_hours} ч. Срок: {course.duration_days} учебных дней. Форма: {course.education_form}. Выдаваемый документ: {course.issued_document}.")
        _paragraph(doc, "1.5. Планируемые результаты обучения", 2)
        _competencies(doc, course)
        _paragraph(doc, "1.6. Календарный учебный график", 2)
        _calendar(doc, course)
        _section(doc, patterns[1])
        _paragraph(doc, "2. Содержание программы", 1)
        _paragraph(doc, "2.1. Учебный план", 2)
        _hours(doc, course, True)
        _section(doc, patterns[-1])
        _paragraph(doc, "2.2. Учебно-тематический план", 2)
        _working_program(doc, course)
        _paragraph(doc, "2.3. Оценочные материалы", 2)
        _assessment(doc, course)
        _paragraph(doc, "2.4. Методические материалы", 2)
        _materials(doc, course)
    _conditions(doc, course, course.template_type == "doop")
    _appendices(doc, course)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def export_pptx(record: ProjectRecord) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for item in record.course.lecture_slides:
        # Split long content into additional slides; retain every bullet and all notes.
        lines = []
        for bullet in item.bullets:
            lines.extend(textwrap.wrap(bullet, width=78, break_long_words=True, break_on_hyphens=False) or [""])
        chunks = [lines[index:index+8] for index in range(0, len(lines), 8)] or [[]]
        title_chunks = textwrap.wrap(item.title, width=58) or [""]
        short_title = "\n".join(title_chunks[:3])
        overflow_title = " ".join(title_chunks[3:])
        if overflow_title:
            extra_lines = textwrap.wrap("Продолжение заголовка: " + overflow_title, width=78)
            lines = extra_lines + lines
            chunks = [lines[index:index+8] for index in range(0, len(lines), 8)]
        for index, chunk in enumerate(chunks):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(.6), Inches(.35), Inches(12.1), Inches(1.5)).text_frame
            title.word_wrap = True
            paragraph = title.paragraphs[0]
            paragraph.text = short_title
            paragraph.font.size = SlidePt(24 if len(title_chunks) > 2 else 28)
            paragraph.font.bold = True
            body = slide.shapes.add_textbox(Inches(.65), Inches(2.05), Inches(12), Inches(4.7)).text_frame
            body.word_wrap = True
            for number, line in enumerate(chunk):
                paragraph = body.paragraphs[0] if number == 0 else body.add_paragraph()
                paragraph.text = line
                paragraph.font.size = SlidePt(22)
                paragraph.space_after = SlidePt(8)
            mode = "Демонстрационный комплект.\n" if record.execution_mode == "demo" else ""
            module = record.slide_modules.get(str(item.slide_number), "")
            if mode:
                footer = slide.shapes.add_textbox(Inches(.65), Inches(7.0), Inches(12), Inches(.25)).text_frame
                footer.paragraphs[0].text = "Демонстрационный комплект"
                footer.paragraphs[0].font.size = SlidePt(12)
            slide.notes_slide.notes_text_frame.text = f"{mode}Модуль: {module}. Исходный слайд: {item.slide_number}.\nПолный заголовок: {item.title}\n{item.speaker_notes}"
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def gift_escape(value: str) -> str:
    value = value.replace("\\", "\\\\")
    value = re.sub(r"([~=#{}:])", r"\\\1", value)
    return value.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\\n")


def export_gift(record: ProjectRecord) -> bytes:
    questions = []
    for question in record.course.tests:
        lines = [f"// Тема: {question.linked_topic_id}; компетенция: {question.linked_competency_id}", f"::{gift_escape(question.id)}::[plain]{gift_escape(question.question)} {{"]
        for option in question.options:
            prefix = "=" if option == question.correct_answer else "~"
            lines.append(f"{prefix}{gift_escape(option)}")
        lines.append("}")
        questions.append("\n".join(lines))
    prefix = "// Демонстрационный комплект для проверки интеграции\n" if record.execution_mode == "demo" else ""
    return (prefix + "\n\n".join(questions) + "\n").encode("utf-8")


def export_csv(record: ProjectRecord) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["id", "question", "options_json", "correct_answer", "topic_id", "competency_id"])
    import json
    for question in record.course.tests:
        writer.writerow([question.id, question.question, json.dumps(question.options, ensure_ascii=False), question.correct_answer, question.linked_topic_id, question.linked_competency_id])
    return output.getvalue().encode("utf-8-sig")
