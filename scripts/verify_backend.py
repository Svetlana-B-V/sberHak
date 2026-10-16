"""Start a private demo server, exercise the checklist and inspect exported files.

Run from the repository root: python scripts/verify_backend.py
No requests to paid models are made, even if the caller has Cloud.ru settings.
"""

import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from uuid import uuid4

from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


ROOT = Path(__file__).resolve().parents[1]


def reference_files():
    document = Document()
    document.add_paragraph("Reference paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "REFERENCE_TABLE_246"
    docx = BytesIO()
    document.save(docx)
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 750 Td (REFERENCE_PDF_123) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf = BytesIO()
    writer.write(pdf)
    return [("reference.docx", docx.getvalue()), ("reference.pdf", pdf.getvalue())]


def multipart(query):
    boundary = "eduagent-" + uuid4().hex
    chunks = [f'--{boundary}\r\nContent-Disposition: form-data; name="query"\r\n\r\n{query}\r\n'.encode()]
    for name, content in reference_files():
        chunks.extend([f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode(), content, b"\r\n"])
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data" / "verification")
    args = parser.parse_args()
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    base = f"http://127.0.0.1:{port}"

    def request(path, body=None, method=None, headers=None, expected=200):
        if isinstance(body, dict):
            body = json.dumps(body, ensure_ascii=False).encode("utf-8")
            headers = {**(headers or {}), "Content-Type": "application/json"}
        item = Request(base + path, data=body, method=method, headers=headers or {})
        try:
            response = urlopen(item, timeout=15)
        except HTTPError as error:
            if error.code == expected:
                return error.read()
            raise RuntimeError(f"{path}: HTTP {error.code}: {error.read().decode('utf-8')}") from error
        with response:
            assert response.status == expected, path
            return response.read()

    def call(path, body=None, method=None):
        return json.loads(request(path, body, method))

    environment = {**os.environ, "ML_MODE": "demo", "CLOUD_RU_API_KEY": "", "PROJECT_DATA_DIR": str(output / "projects"), "PYTHONIOENCODING": "utf-8"}
    with (output / "server.log").open("w", encoding="utf-8") as server_log:
        server = subprocess.Popen(
            [sys.executable, "-m", "uvicorn", "backend.main:app", "--host", "127.0.0.1", "--port", str(port)],
            cwd=ROOT, env=environment, stdin=subprocess.DEVNULL, stdout=server_log, stderr=subprocess.STDOUT,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        try:
            deadline = time.monotonic() + 20
            while time.monotonic() < deadline:
                if server.poll() is not None:
                    raise RuntimeError((output / "server.log").read_text(encoding="utf-8"))
                try:
                    with urlopen(base + "/api/health", timeout=1):
                        break
                except (URLError, OSError):
                    time.sleep(.1)
            else:
                raise RuntimeError("The demo server did not start")
            assert b"swagger-ui" in request("/docs")
            assert "/api/projects/skeleton" in call("/openapi.json")["paths"]
            results = []
            for template, hours in [("dpp_pk", 72), ("doop", 16)]:
                query = f"{'Повышение квалификации' if template == 'dpp_pk' else 'Общеразвивающая программа'} по Python для сотрудников, {hours} часов, очная форма, итоговая аттестация — защита проекта."
                data, content_type = multipart(query)
                analysis = json.loads(request("/api/projects/analyze", data, "POST", {"Content-Type": content_type}))
                project_id = analysis["project_id"]
                course = call("/api/projects/skeleton", {"project_id": project_id, "query": query, "template_type": template, "total_hours": hours})
                path = f"/api/projects/{project_id}"
                assert len(course["reference_files"]) == 2
                request(path + "/generate-content", b"", "POST", expected=409)
                course = call(path + "/skeleton", course, "PUT")
                course = call(path + "/generate-content", {}, "POST")
                assert course["status"] == "completed" and course["execution_mode"] == "demo"
                assert course["validation_report"]["metrics"]["api_calls"] == 0
                assert sum(mod["total_hours"] for mod in course["modules"]) + course["ia_hours"] == hours
                directory = output / template
                directory.mkdir(exist_ok=True)
                files = {}
                for extension in ("docx", "pptx", "gift", "csv", "json"):
                    content = request(path + "/export/" + extension)
                    destination = directory / ("course." + extension)
                    destination.write_bytes(content)
                    files[extension] = len(content)
                    if extension == "docx":
                        assert Document(BytesIO(content)).tables
                    elif extension == "pptx":
                        assert len(Presentation(BytesIO(content)).slides) > 0
                before = course["modules"][1]
                course = call(path + "/regenerate-part", {"element_type": "topic", "element_id": "1.1", "instruction": "Добавить пример проверки данных", "revision": course["revision"]})
                assert course["modules"][1] == before
                assert call(path)["revision"] == course["revision"]
                record = json.loads((output / "projects" / f"{project_id}.json").read_text(encoding="utf-8"))
                reference_text = "\n".join(item["text"] for item in record["references"])
                assert "REFERENCE_TABLE_246" in reference_text and "REFERENCE_PDF_123" in reference_text
                results.append({"template": template, "total_hours": hours, "file_sizes": files, "status": "passed"})
            report = {"execution_mode": "demo", "real_cloud_tested": False, "visual_office_tested": False, "python_version": sys.version.split()[0], "cases": results}
            (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps(report, ensure_ascii=False, indent=2))
        finally:
            server.terminate()
            try:
                server.wait(timeout=5)
            except subprocess.TimeoutExpired:
                server.kill()
                server.wait(timeout=5)


if __name__ == "__main__":
    main()
