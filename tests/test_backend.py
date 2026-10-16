from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from threading import Event

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest

from backend.config import Settings
from backend.errors import BackendError
from backend.exporters import gift_escape
from backend.main import create_app
from backend.skills import Skills, Session, SkillClient, _session


QUERY = "Программа повышения квалификации по анализу данных на Python для сотрудников банков, 72 часа, очно-заочная форма, итоговая аттестация — защита проекта."


@pytest.fixture
def app(tmp_path):
    return create_app(Settings(data_dir=tmp_path / "projects", ml_mode="demo"))


@pytest.fixture
def client(app):
    with TestClient(app) as connection:
        yield connection


def skeleton(client, template="dpp_pk", hours=72):
    response = client.post("/api/projects/skeleton", json={"query": QUERY, "template_type": template, "total_hours": hours})
    assert response.status_code == 200, response.text
    return response.json()


def completed(client, template="dpp_pk", hours=72):
    course = skeleton(client, template, hours)
    project_id = course["project_id"]
    response = client.put(f"/api/projects/{project_id}/skeleton", json=course)
    assert response.status_code == 200, response.text
    response = client.post(f"/api/projects/{project_id}/generate-content")
    assert response.status_code == 200, response.text
    return response.json()


def docx_reference():
    document = Document()
    document.add_paragraph("REFERENCE_PARAGRAPH")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "REFERENCE_TABLE_246"
    table.cell(0, 1).text = "СРС"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "NESTED_REFERENCE"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_reference(text=True, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 750 Td (REFERENCE_PDF_123) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("template,hours", [("dpp_pk", 72), ("doop", 16)])
def test_full_workflow_and_exports(client, app, template, hours):
    course = completed(client, template, hours)
    project_id = course["project_id"]
    assert course["status"] == "completed"
    assert course["execution_mode"] == "demo"
    assert course["validation_report"]["metrics"]["api_calls"] == 0
    assert course["validation_report"]["metrics"]["estimated_cost_rub"] is None
    assert sum(module["total_hours"] for module in course["modules"]) + course["ia_hours"] == hours
    saved = json.loads((app.state.settings.data_dir / f"{project_id}.json").read_text(encoding="utf-8"))
    assert saved["course"]["title"] == course["title"]
    if template == "doop":
        assert all(not comp["profstandard_code"] and not comp["labor_function_code"] for comp in course["competencies"])
    for extension in ["docx", "pptx", "gift", "csv", "json"]:
        response = client.get(f"/api/projects/{project_id}/export/{extension}")
        assert response.status_code == 200, response.text
        assert "attachment" in response.headers["content-disposition"]
        if extension == "docx":
            document = Document(BytesIO(response.content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            assert course["title"] in text
            assert 'Приказом № ______ от «___» __________ 20__ г.' in text
            assert "17-ОБР" not in text
            assert "Кадровое обеспечение" in text
            assert "Тестовые вопросы" in text
            expected_sections = 3 if template == "doop" else 1
            assert len(document.sections) == expected_sections
            for table in document.tables:
                if table.rows[0].cells[0].text == "№":
                    assert table.rows[0].cells[7].text == "ИА"
                    totals = table.rows[-1].cells
                    assert sum(int(totals[i].text) for i in range(3, 8)) == int(totals[2].text)
        elif extension == "pptx":
            deck = Presentation(BytesIO(response.content))
            assert len(deck.slides) >= len(course["lecture_slides"])
            assert " ".join(deck.slides[0].shapes[0].text.split()) == course["lecture_slides"][0]["title"]
            assert any("Модуль:" in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        elif extension == "gift":
            text = response.content.decode("utf-8")
            assert text.count("::[plain]") == len(course["tests"])
            assert "}\n\n//" in text
        elif extension == "csv":
            assert "options_json" in response.content.decode("utf-8-sig")
        else:
            assert response.json()["project_id"] == project_id


@pytest.mark.parametrize("hours", [16, 36, 72, 248, 256, 500])
def test_hours_match_at_supported_sizes(client, hours):
    course = skeleton(client, "doop" if hours == 16 else "dpp_pk", hours)
    assert sum(module["total_hours"] for module in course["modules"]) + course["ia_hours"] == hours
    for module in course["modules"]:
        for topic in module["topics"]:
            assert all(topic[name] >= 0 for name in ("lecture_hours", "practice_hours", "lab_hours", "srs_hours"))


def test_approval_gate_and_stale_materials(client):
    course = skeleton(client)
    url = f"/api/projects/{course['project_id']}"
    assert client.post(url + "/generate-content").status_code == 409
    assert client.get(url + "/export/docx").status_code == 409
    course = completed(client)
    url = f"/api/projects/{course['project_id']}"
    response = client.put(url + "/skeleton", json={"goal": "Новая согласованная цель", "revision": course["revision"]})
    assert response.status_code == 200
    assert response.json()["tests"] == []
    assert response.json()["status"] == "approved"
    assert client.get(url + "/export/docx").status_code == 409


def test_invalid_or_stale_updates_do_not_overwrite_project(client):
    course = skeleton(client)
    url = f"/api/projects/{course['project_id']}"
    assert client.put(url + "/skeleton", json={"project_id": "another-project"}).status_code == 422
    bad = json.loads(json.dumps(course))
    bad["modules"][0]["topics"][0]["srs_hours"] = -5
    assert client.put(url + "/skeleton", json=bad).status_code == 422
    assert client.get(url).json()["revision"] == course["revision"]
    approved = client.put(url + "/skeleton", json=course)
    assert approved.status_code == 200
    assert client.put(url + "/skeleton", json=course).status_code == 409


@pytest.mark.parametrize("kind", ["docx", "pdf"])
def test_references_survive_analysis_skeleton_and_restart(client, app, kind):
    data = docx_reference() if kind == "docx" else pdf_reference()
    response = client.post("/api/projects/analyze", data={"query": QUERY}, files=[("files", (f"../../reference.{kind}", data))])
    assert response.status_code == 200, response.text
    project_id = response.json()["project_id"]
    response = client.post("/api/projects/skeleton", json={"query": QUERY, "project_id": project_id})
    assert response.status_code == 200, response.text
    saved = app.state.service.store.load(project_id)
    if kind == "docx":
        assert "REFERENCE_TABLE_246" in saved.reference_text()
        assert "NESTED_REFERENCE" in saved.reference_text()
    else:
        assert "REFERENCE_PDF_123" in saved.reference_text()
    assert saved.references[0].filename == f"reference.{kind}"
    second_app = create_app(app.state.settings)
    with TestClient(second_app) as second_client:
        response = second_client.get(f"/api/projects/{project_id}")
        assert response.status_code == 200
        assert response.json()["reference_files"][0]["sha256"] == saved.references[0].sha256


@pytest.mark.parametrize("name,data,code", [
    ("bad.docx", b"not a zip", 422),
    ("bad.exe", b"MZ", 415),
    ("empty.pdf", b"", 422),
])
def test_invalid_references(client, name, data, code):
    response = client.post("/api/projects/analyze", data={"query": QUERY}, files={"files": (name, data)})
    assert response.status_code == code


@pytest.mark.parametrize("text,encrypted", [(False, False), (True, True)])
def test_scanned_or_encrypted_pdf_is_not_silently_ignored(client, text, encrypted):
    response = client.post("/api/projects/analyze", data={"query": QUERY}, files={"files": ("reference.pdf", pdf_reference(text, encrypted))})
    assert response.status_code == 422


def test_upload_limits(app, tmp_path):
    settings = replace(app.state.settings, data_dir=tmp_path / "limits", max_file_bytes=100)
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/projects/analyze", data={"query": QUERY}, files={"files": ("reference.docx", docx_reference())})
        assert response.status_code == 413


def test_topic_regeneration_refreshes_dependencies_and_preserves_other_module(client):
    original = completed(client)
    url = f"/api/projects/{original['project_id']}"
    response = client.post(url + "/regenerate-part", json={"element_type": "topic", "element_id": "1.1", "instruction": "Добавить пример по проверке пропущенных значений", "revision": original["revision"]})
    assert response.status_code == 200, response.text
    changed = response.json()
    assert changed["modules"][1] == original["modules"][1]
    assert changed["modules"][0]["topics"][1:] == original["modules"][0]["topics"][1:]
    assert changed["modules"][0]["topics"][0]["total_hours"] == original["modules"][0]["topics"][0]["total_hours"]
    assert changed["modules"][0]["topics"][0]["title"] != original["modules"][0]["topics"][0]["title"]
    assert [q for q in changed["tests"] if q["linked_topic_id"] != "1.1"] == [q for q in original["tests"] if q["linked_topic_id"] != "1.1"]
    assert [s for s in changed["seminars"] if s["linked_topic_id"] != "1.1"] == [s for s in original["seminars"] if s["linked_topic_id"] != "1.1"]
    assert [s for s in changed["lecture_slides"] if s["module_id"] == "2"] == [s for s in original["lecture_slides"] if s["module_id"] == "2"]
    assert any("обновлено" in q["question"] for q in changed["tests"] if q["linked_topic_id"] == "1.1")


def test_test_regeneration_changes_only_selected_question(client):
    original = completed(client)
    target_id = original["tests"][0]["id"]
    url = f"/api/projects/{original['project_id']}"
    response = client.post(url + "/regenerate-part", json={"element_type": "test", "element_id": target_id, "instruction": "Усложнить вопрос"})
    assert response.status_code == 200, response.text
    changed = response.json()
    assert changed["modules"] == original["modules"]
    assert changed["lecture_slides"] == original["lecture_slides"]
    assert changed["tests"][1:] == original["tests"][1:]
    assert changed["tests"][0]["question"] != original["tests"][0]["question"]
    assert client.post(url + "/regenerate-part", json={"element_type": "topic", "element_id": "missing", "instruction": "Изменить"}).status_code == 404


def test_unified_endpoint_preserves_human_approval(client):
    response = client.post("/api/v1/evaluate", json={"query": "Курс по кибербезопасности для студентов колледжа примерно на месяц"})
    assert response.status_code == 200
    assert response.json()["status"] == "clarifying"
    assert 1 <= len(response.json()["clarifying_questions"]) <= 5
    assert "total_hours" not in response.json()["extracted_params"]
    response = client.post("/api/v1/evaluate", json={"query": QUERY})
    assert response.status_code == 200, response.text
    project_id = response.json()["project_id"]
    assert response.json()["status"] == "skeleton_ready"
    assert client.post("/api/v1/evaluate", json={"action": "generate", "project_id": project_id}).status_code == 409
    response = client.post("/api/v1/evaluate", json={"action": "approve", "project_id": project_id, "approved": True})
    assert response.status_code == 200, response.text
    response = client.post("/api/v1/evaluate", json={"action": "generate", "project_id": project_id})
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "completed"


def test_cloud_mode_without_credentials_returns_error_instead_of_demo(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path / "cloud", ml_mode="cloud"))) as client:
        response = client.post("/api/projects/skeleton", json={"query": QUERY})
        assert response.status_code == 503
        assert response.json()["detail"]["code"] == "cloud_not_configured"


def test_provider_failure_never_returns_fallback():
    class Failure:
        def create(self, **kwargs):
            raise RuntimeError("provider unavailable")
    class FakeClient:
        chat = type("Chat", (), {"completions": Failure()})()
    session = Session(Settings(ml_mode="cloud"), FakeClient())
    token = _session.set(session)
    try:
        with pytest.raises(BackendError) as error:
            SkillClient().generate_json("system", "user", {"fake": True})
        assert error.value.status_code == 502
    finally:
        _session.reset(token)


def test_generate_failure_preserves_approved_skeleton(app):
    real = app.state.service.skills
    class FailureSkills:
        def run(self, operation, *args, **kwargs):
            if operation == "content":
                raise BackendError(502, "test_provider_failure", "Ошибка поставщика")
            return real.run(operation, *args, **kwargs)
    app.state.service.skills = FailureSkills()
    with TestClient(app) as client:
        course = skeleton(client)
        url = f"/api/projects/{course['project_id']}"
        approved = client.put(url + "/skeleton", json=course).json()
        assert client.post(url + "/generate-content").status_code == 502
        saved = client.get(url).json()
        assert saved["status"] == "approved"
        assert saved["busy"] is None
        assert saved["modules"] == approved["modules"]


def test_health_remains_responsive_during_ml_call(app):
    real = app.state.service.skills
    started, release = Event(), Event()
    class SlowSkills:
        def run(self, operation, *args, **kwargs):
            if operation == "content":
                started.set()
                if not release.wait(4):
                    raise RuntimeError("Test did not release the ML call")
            return real.run(operation, *args, **kwargs)
    app.state.service.skills = SlowSkills()
    with TestClient(app) as client:
        course = skeleton(client)
        url = f"/api/projects/{course['project_id']}"
        client.put(url + "/skeleton", json=course)
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(client.post, url + "/generate-content")
            try:
                assert started.wait(3)
                assert client.get("/api/health").status_code == 200
                assert client.get(url).json()["busy"] == "generating"
            finally:
                release.set()
            assert future.result(timeout=5).status_code == 200


def test_gift_reserved_characters_are_escaped():
    assert gift_escape("x={a}~b#c:d\\e\nf") == "x\\=\\{a\\}\\~b\\#c\\:d\\\\e\\nf"
